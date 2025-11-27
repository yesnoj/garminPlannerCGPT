import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import pandas as pd
import json
from datetime import datetime
import threading

from excel_utils import generate_training_excel, format_workbook_dates_and_steps
from dsl_parser import expand_repeat_lines, build_garmin_workout_from_excel_row
from garmin_service import GarminService


class LoadingDialog:
    """Finestra di loading modale con animazione."""
    
    def __init__(self, parent, title="Operazione in corso..."):
        self.top = tk.Toplevel(parent)
        self.top.title(title)
        self.top.transient(parent)
        self.top.grab_set()
        
        # Centra la finestra
        self.top.geometry("300x100")
        self.top.resizable(False, False)
        
        # Frame principale
        frame = ttk.Frame(self.top, padding=20)
        frame.pack(fill=tk.BOTH, expand=True)
        
        # Label con messaggio
        self.label = ttk.Label(frame, text="Attendere...", font=("", 10))
        self.label.pack(pady=(0, 10))
        
        # Progress bar indeterminata
        self.progress = ttk.Progressbar(frame, mode='indeterminate', length=250)
        self.progress.pack()
        self.progress.start(10)
        
        # Impedisci chiusura con X
        self.top.protocol("WM_DELETE_WINDOW", lambda: None)
        
        # Centra rispetto al parent
        self.top.update_idletasks()
        x = parent.winfo_x() + (parent.winfo_width() // 2) - (self.top.winfo_width() // 2)
        y = parent.winfo_y() + (parent.winfo_height() // 2) - (self.top.winfo_height() // 2)
        self.top.geometry(f"+{x}+{y}")
    
    def update_message(self, message: str):
        """Aggiorna il messaggio mostrato."""
        self.label.config(text=message)
        self.top.update_idletasks()
    
    def close(self):
        """Chiude la finestra di loading."""
        self.progress.stop()
        self.top.grab_release()
        self.top.destroy()


class TrainingPlannerGUI:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("Trainer Planner (Excel + Garmin)")

        self.excel_path: str | None = None
        self.df_workouts: pd.DataFrame | None = None
        self.df_parameters: pd.DataFrame | None = None

        self.garmin = GarminService()
        self.use_prefix_var = tk.BooleanVar(value=True)


        self._build_ui()

    # ---------- UI ----------

    def _build_ui(self):
        top_frame = ttk.Frame(self.root, padding=5)
        top_frame.pack(side=tk.TOP, fill=tk.X)

        btn_generate = ttk.Button(top_frame, text="Genera Excel di esempio", command=self.generate_example_excel)
        btn_generate.pack(side=tk.LEFT, padx=5)

        btn_load = ttk.Button(top_frame, text="Carica Excel", command=self.load_excel)
        btn_load.pack(side=tk.LEFT, padx=5)

        btn_save = ttk.Button(top_frame, text="Salva Excel", command=self.save_excel)
        btn_save.pack(side=tk.LEFT, padx=5)

        self.lbl_file = ttk.Label(top_frame, text="Nessun file caricato")
        self.lbl_file.pack(side=tk.LEFT, padx=10)

        # Garmin frame
        garmin_frame = ttk.LabelFrame(self.root, text="Garmin Connect", padding=10)
        garmin_frame.pack(side=tk.TOP, fill=tk.X, padx=5, pady=5)

        # Configurazione colonne per layout responsive
        garmin_frame.columnconfigure(1, weight=1)  # Entry email/password si espandono

        # ========== SEZIONE AUTENTICAZIONE ==========
        
        # Riga 0: Email
        ttk.Label(garmin_frame, text="Email:").grid(
            row=0, column=0, sticky="e", padx=(0, 8), pady=(0, 5)
        )
        self.entry_email = ttk.Entry(garmin_frame, width=35)
        self.entry_email.grid(
            row=0, column=1, columnspan=2, sticky="we", padx=(0, 10), pady=(0, 5)
        )
        
        # Riga 1: Password
        ttk.Label(garmin_frame, text="Password:").grid(
            row=1, column=0, sticky="e", padx=(0, 8), pady=(0, 8)
        )
        self.entry_password = ttk.Entry(garmin_frame, width=35, show="*")
        self.entry_password.grid(
            row=1, column=1, columnspan=2, sticky="we", padx=(0, 10), pady=(0, 8)
        )
        
        # Bottoni login sulla destra
        btn_frame_login = ttk.Frame(garmin_frame)
        btn_frame_login.grid(row=0, column=3, rowspan=2, sticky="nsew", padx=(10, 0))
        
        btn_login_creds = ttk.Button(
            btn_frame_login,
            text="Login (crea/aggiorna sessione)",
            command=self.login_garmin_with_credentials,
        )
        btn_login_creds.pack(side=tk.TOP, fill=tk.X, pady=(0, 5))

        btn_login_session = ttk.Button(
            btn_frame_login,
            text="Usa sessione salvata",
            command=self.login_garmin_with_session,
        )
        btn_login_session.pack(side=tk.TOP, fill=tk.X)
        
        # Status sulla destra dei bottoni login
        self.lbl_garmin_status = ttk.Label(
            garmin_frame, text="● Non connesso", foreground="red", font=("", 9, "bold")
        )
        self.lbl_garmin_status.grid(
            row=0, column=4, rowspan=2, sticky="w", padx=(15, 0)
        )

        # Separatore
        ttk.Separator(garmin_frame, orient="horizontal").grid(
            row=2, column=0, columnspan=5, sticky="ew", pady=10
        )

        # ========== SEZIONE OPZIONI ==========
        
        # Riga 3: Checkbox prefisso
        chk_prefix = ttk.Checkbutton(
            garmin_frame,
            text="☑ Prefisso nome WnSn (Week/Session)",
            variable=self.use_prefix_var,
        )
        chk_prefix.grid(
            row=3, column=0, columnspan=5, sticky="w", pady=(0, 10)
        )

        # Separatore
        ttk.Separator(garmin_frame, orient="horizontal").grid(
            row=4, column=0, columnspan=5, sticky="ew", pady=(0, 10)
        )

        # ========== SEZIONE OPERAZIONI WORKOUT ==========
        
        # Frame per bottoni operazioni (riga 5)
        btn_operations_frame = ttk.Frame(garmin_frame)
        btn_operations_frame.grid(row=5, column=0, columnspan=5, sticky="ew")
        
        # Configura le colonne del frame operazioni per distribuire uniformemente
        for i in range(4):
            btn_operations_frame.columnconfigure(i, weight=1)
        
        # Bottoni operazioni principali
        btn_upload = ttk.Button(
            btn_operations_frame,
            text="📤 Carica su Garmin",
            command=self.upload_selected_workouts,
        )
        btn_upload.grid(row=0, column=0, sticky="ew", padx=(0, 5))

        btn_upload_plan = ttk.Button(
            btn_operations_frame,
            text="📤📅 Carica + Pianifica",
            command=self.upload_and_schedule_selected_workouts,
        )
        btn_upload_plan.grid(row=0, column=1, sticky="ew", padx=(0, 5))

        btn_unschedule = ttk.Button(
            btn_operations_frame,
            text="📅✖ Rimuovi Pianificazione",
            command=self.unschedule_selected_workouts,
        )
        btn_unschedule.grid(row=0, column=2, sticky="ew", padx=(0, 5))

        btn_delete = ttk.Button(
            btn_operations_frame,
            text="🗑 Cancella da Garmin",
            command=self.delete_workouts_from_garmin,
        )
        btn_delete.grid(row=0, column=3, sticky="ew")


        # Main split
        main_pane = ttk.Panedwindow(self.root, orient=tk.HORIZONTAL)
        main_pane.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)

        # Left: workouts
        left_frame = ttk.Frame(main_pane, padding=5)
        main_pane.add(left_frame, weight=1)

        ttk.Label(left_frame, text="Workouts").pack(anchor="w")

        cols = ("Week", "Date", "Session", "Sport", "Description", "ScheduledDate")
        self.tree_workouts = ttk.Treeview(
            left_frame,
            columns=cols,
            show="headings",
            height=15,
            selectmode="extended",
        )
        headers = {
            "Week": "Week",
            "Date": "Date",
            "Session": "Session",
            "Sport": "Sport",
            "Description": "Description",
            "ScheduledDate": "Sched.",
        }
        widths = {
            "Week": 50,
            "Date": 90,
            "Session": 70,
            "Sport": 80,
            "Description": 220,
            "ScheduledDate": 90,
        }
        for c in cols:
            self.tree_workouts.heading(c, text=headers[c])
            self.tree_workouts.column(c, width=widths[c], anchor="w")

        # Tag colori per stato workout
        self.tree_workouts.tag_configure("scheduled", background="#d9fdd3")      # verde
        self.tree_workouts.tag_configure("uploaded", background="#fff4ce")       # giallo
        self.tree_workouts.tag_configure("not_uploaded", background="#ffffff")   # bianco

        self.tree_workouts.pack(fill=tk.BOTH, expand=True)
        self.tree_workouts.bind("<<TreeviewSelect>>", self.on_select_workout)

        # Right
        right_frame = ttk.Frame(main_pane, padding=5)
        main_pane.add(right_frame, weight=1)

        notebook = ttk.Notebook(right_frame)
        notebook.pack(fill=tk.BOTH, expand=True)

        # Tab Steps
        steps_frame = ttk.Frame(notebook, padding=5)
        notebook.add(steps_frame, text="Steps")

        meta_frame = ttk.Frame(steps_frame)
        meta_frame.pack(fill=tk.X, pady=2)

        ttk.Label(meta_frame, text="Week:").grid(row=0, column=0, sticky="w")
        self.entry_week = ttk.Entry(meta_frame, width=5)
        self.entry_week.grid(row=0, column=1, sticky="w", padx=(0, 10))

        ttk.Label(meta_frame, text="Session:").grid(row=0, column=2, sticky="w")
        self.entry_session = ttk.Entry(meta_frame, width=5)
        self.entry_session.grid(row=0, column=3, sticky="w", padx=(0, 10))

        ttk.Label(meta_frame, text="Date (YYYY-MM-DD):").grid(row=0, column=4, sticky="w")
        self.entry_date = ttk.Entry(meta_frame, width=12)
        self.entry_date.grid(row=0, column=5, sticky="w", padx=(0, 10))

        ttk.Label(steps_frame, text="Description:").pack(anchor="w")
        self.entry_description = ttk.Entry(steps_frame)
        self.entry_description.pack(fill=tk.X, pady=2)

        ttk.Label(steps_frame, text="Steps (testo):").pack(anchor="w")
        self.text_steps = tk.Text(steps_frame, wrap="none", height=10)
        self.text_steps.pack(fill=tk.BOTH, expand=True)

        btn_frame = ttk.Frame(steps_frame)
        btn_frame.pack(fill=tk.X, pady=3)

        btn_save_workout = ttk.Button(btn_frame, text="Salva modifiche workout", command=self.update_workout_row)
        btn_save_workout.pack(side=tk.LEFT)

        btn_show_json = ttk.Button(btn_frame, text="Mostra JSON steps espanso", command=self.show_steps_json)
        btn_show_json.pack(side=tk.LEFT, padx=5)

        self.text_json = tk.Text(steps_frame, wrap="none", height=10)
        self.text_json.pack(fill=tk.BOTH, expand=True)

        # Tab Parameters
        params_frame = ttk.Frame(notebook, padding=5)
        notebook.add(params_frame, text="Parameters")

        self.tree_params = ttk.Treeview(
            params_frame,
            columns=("Key", "Metric", "Expression", "Notes"),
            show="headings",
            height=15,
            selectmode="browse",
        )
        self.tree_params.heading("Key", text="Key")
        self.tree_params.heading("Metric", text="Metric")
        self.tree_params.heading("Expression", text="Expression")
        self.tree_params.heading("Notes", text="Notes")

        self.tree_params.column("Key", width=120, anchor="w")
        self.tree_params.column("Metric", width=70, anchor="w")
        self.tree_params.column("Expression", width=120, anchor="w")
        self.tree_params.column("Notes", width=220, anchor="w")

        self.tree_params.pack(fill=tk.BOTH, expand=True)

        params_btn_frame = ttk.Frame(params_frame)
        params_btn_frame.pack(fill=tk.X, pady=5)

        btn_edit_param = ttk.Button(params_btn_frame, text="Modifica parametro", command=self.edit_selected_parameter)
        btn_edit_param.pack(side=tk.LEFT)

    # ---------- Excel ----------

    def generate_example_excel(self):
        path = filedialog.asksaveasfilename(
            title="Salva Excel di esempio",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not path:
            return
        try:
            generate_training_excel(path)
            messagebox.showinfo("OK", f"File di esempio creato:\n{path}")
        except Exception as e:
            messagebox.showerror("Errore", f"Errore nella generazione del file:\n{e}")

    def load_excel(self):
        path = filedialog.askopenfilename(
            title="Seleziona file Excel",
            filetypes=[("Excel", "*.xlsx *.xls")],
        )
        if not path:
            return
        try:
            self.df_workouts = pd.read_excel(path, sheet_name="Workouts")

            try:
                # Ã°Å¸â€˜â€° forza tutte le colonne di Parameters a stringa
                self.df_parameters = pd.read_excel(
                    path,
                    sheet_name="Parameters",
                    dtype={"Key": str, "Metric": str, "Expression": str, "Notes": str},
                )
            except Exception:
                self.df_parameters = pd.DataFrame(columns=["Key", "Metric", "Expression", "Notes"])

            for col in ["WorkoutId", "ScheduledDate"]:
                if col not in self.df_workouts.columns:
                    self.df_workouts[col] = ""

            # forza tipo stringa per evitare future warning
            if "WorkoutId" in self.df_workouts.columns:
                self.df_workouts["WorkoutId"] = self.df_workouts["WorkoutId"].astype("string")
            if "ScheduledDate" in self.df_workouts.columns:
                self.df_workouts["ScheduledDate"] = self.df_workouts["ScheduledDate"].astype("string")

            self.excel_path = path
            self.lbl_file.config(text=path)
            self.populate_workouts_tree()
            self.populate_params_tree()

            self.entry_week.delete(0, tk.END)
            self.entry_date.delete(0, tk.END)
            self.entry_session.delete(0, tk.END)
            self.entry_description.delete(0, tk.END)
            self.text_steps.delete("1.0", tk.END)
            self.text_json.delete("1.0", tk.END)
        except Exception as e:
            messagebox.showerror("Errore", f"Impossibile caricare l'Excel:\n{e}")



    def write_excel(self, path: str):
        if self.df_workouts is None:
            return

        # Copia temporanea per non modificare l'originale
        df_temp = self.df_workouts.copy()

        # Date -> datetime
        if "Date" in df_temp.columns:
            df_temp["Date"] = pd.to_datetime(
                df_temp["Date"], errors="coerce", dayfirst=True
            )

        # ScheduledDate -> converti in date (solo data, no orario)
        if "ScheduledDate" in df_temp.columns:
            def to_date_only(val):
                if pd.isna(val) or val == "" or str(val).strip().lower() in ("nan", "<na>"):
                    return None
                try:
                    dt = pd.to_datetime(val, errors="coerce")
                    return dt.date() if pd.notna(dt) else None
                except:
                    return None
            
            df_temp["ScheduledDate"] = df_temp["ScheduledDate"].apply(to_date_only)

        # Carica eventuali fogli esistenti da preservare
        existing_sheets = {}
        try:
            from openpyxl import load_workbook
            wb_existing = load_workbook(path)
            
            # Salva il foglio "Esempi DSL" se esiste
            if "Esempi DSL" in wb_existing.sheetnames:
                df_esempi = pd.read_excel(path, sheet_name="Esempi DSL")
                existing_sheets["Esempi DSL"] = df_esempi
                
            wb_existing.close()
        except:
            pass

        # Scrivi l'Excel con tutti i fogli
        with pd.ExcelWriter(path, engine="openpyxl", mode='w') as writer:
            df_temp.to_excel(writer, sheet_name="Workouts", index=False)
            
            if self.df_parameters is not None:
                self.df_parameters.to_excel(writer, sheet_name="Parameters", index=False)
            
            # Riscri gli fogli preservati
            for sheet_name, df in existing_sheets.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False)

        # Applica formattazione finale (dd/mm/yyyy, Steps multilinea, ecc.)
        format_workbook_dates_and_steps(path)


    def save_excel(self):
        if self.df_workouts is None:
            messagebox.showerror("Errore", "Nessun allenamento caricato.")
            return

        path = filedialog.asksaveasfilename(
            title="Salva Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not path:
            return
        try:
            self.write_excel(path)
            messagebox.showinfo("OK", f"File salvato:\n{path}")
        except Exception as e:
            messagebox.showerror("Errore", f"Errore nel salvataggio:\n{e}")

    def autosave_to_loaded_excel(self):
        """Salva automaticamente sul file Excel da cui abbiamo caricato (se possibile)."""
        if not self.excel_path:
            return
        if self.df_workouts is None:
            return
        try:
            self.write_excel(self.excel_path)
        except Exception as e:
            messagebox.showwarning(
                "Autosalvataggio fallito",
                f"Non riesco a salvare su:\n{self.excel_path}\n\nDettagli:\n{e}\n\n"
                "Chiudi il file in Excel (se ÃƒÆ’Ã‚Â¨ aperto) e prova a salvare manualmente."
            )

    def populate_workouts_tree(self):
        for item in self.tree_workouts.get_children():
            self.tree_workouts.delete(item)
        if self.df_workouts is None:
            return
        for idx, row in self.df_workouts.iterrows():
            week = row.get("Week", "")
            date = self.format_date_for_display(row.get("Date", ""))
            session = row.get("Session", "")
            sport = row.get("Sport", "")
            desc = row.get("Description", "")
            sched = row.get("ScheduledDate", "")

            if pd.isna(sched) or str(sched).strip().lower() in ("nan", "<na>"):
                sched = ""

            vals = (week, date, session, sport, desc, sched)
            iid = str(idx)
            self.tree_workouts.insert("", tk.END, iid=iid, values=vals)
            self.apply_workout_row_style(idx)


    def format_date_for_display(self, value) -> str:
        """Converte il valore della colonna Date in una stringa YYYY-MM-DD per la GUI."""
        import pandas as pd
        from datetime import datetime, date as date_cls

        if value is None or (isinstance(value, float) and pd.isna(value)):
            return ""

        # pandas Timestamp o datetime
        if isinstance(value, (pd.Timestamp, datetime)):
            return value.date().isoformat()

        # oggetto date puro
        if isinstance(value, date_cls):
            return value.isoformat()

        # stringa generica
        s = str(value).strip()
        # caso "2025-11-24 00:00:00" -> prendo solo la parte prima dello spazio
        if " " in s:
            s = s.split(" ")[0]
        return s


    def populate_params_tree(self):
        for item in self.tree_params.get_children():
            self.tree_params.delete(item)
        if self.df_parameters is None or self.df_parameters.empty:
            return
        for idx, row in self.df_parameters.iterrows():
            vals = (
                row.get("Key", ""),
                row.get("Metric", ""),
                row.get("Expression", ""),
                row.get("Notes", ""),
            )
            self.tree_params.insert("", tk.END, iid=str(idx), values=vals)

    # ---------- Workouts editing ----------

    def on_select_workout(self, event=None):
        if self.df_workouts is None:
            return
        sel = self.tree_workouts.selection()
        if not sel:
            return
        idx = int(sel[-1])
        row = self.df_workouts.iloc[idx]

        self.entry_week.delete(0, tk.END)
        self.entry_week.insert(0, str(row.get("Week", "")))

        self.entry_date.delete(0, tk.END)
        self.entry_date.insert(0, self.format_date_for_display(row.get("Date", "")))

        self.entry_session.delete(0, tk.END)
        self.entry_session.insert(0, str(row.get("Session", "")))

        self.entry_description.delete(0, tk.END)
        self.entry_description.insert(0, str(row.get("Description", "")))

        steps_text = str(row.get("Steps", ""))
        self.text_steps.delete("1.0", tk.END)
        self.text_steps.insert("1.0", steps_text)

        self.text_json.delete("1.0", tk.END)

    def update_workout_row(self):
        if self.df_workouts is None:
            return
        sel = self.tree_workouts.selection()
        if not sel:
            messagebox.showwarning("Attenzione", "Seleziona almeno un allenamento.")
            return

        week_str = self.entry_week.get().strip()
        date_val = self.entry_date.get().strip()
        session_str = self.entry_session.get().strip()
        desc = self.entry_description.get().strip()
        steps_text = self.text_steps.get("1.0", tk.END).rstrip("\n")

        for iid in sel:
            idx = int(iid)
            if week_str:
                try:
                    self.df_workouts.at[idx, "Week"] = int(week_str)
                except ValueError:
                    self.df_workouts.at[idx, "Week"] = week_str
            if date_val:
                self.df_workouts.at[idx, "Date"] = date_val
            if session_str:
                try:
                    self.df_workouts.at[idx, "Session"] = int(session_str)
                except ValueError:
                    self.df_workouts.at[idx, "Session"] = session_str
            self.df_workouts.at[idx, "Description"] = desc
            self.df_workouts.at[idx, "Steps"] = steps_text

            self.refresh_tree_row(idx)

        messagebox.showinfo("OK", "Allenamento/i aggiornato/i.")

    def show_steps_json(self):
        steps_text = self.text_steps.get("1.0", tk.END).rstrip("\n")
        flat_lines = expand_repeat_lines(steps_text)
        data = {"flat_lines": flat_lines}
        self.text_json.delete("1.0", tk.END)
        self.text_json.insert("1.0", json.dumps(data, indent=2, ensure_ascii=False))

    def apply_workout_row_style(self, idx: int):
        """Applica il colore giusto alla riga in base a WorkoutId / ScheduledDate."""
        if self.df_workouts is None:
            return

        row = self.df_workouts.iloc[idx]
        iid = str(idx)

        workout_id = str(row.get("WorkoutId", "")).strip()
        sched_date = str(row.get("ScheduledDate", "")).strip()

        if workout_id.lower() in ("nan", "<na>"):
            workout_id = ""
        if sched_date.lower() in ("nan", "<na>"):
            sched_date = ""

        if sched_date:
            tag = "scheduled"
        elif workout_id:
            tag = "uploaded"
        else:
            tag = "not_uploaded"

        self.tree_workouts.item(iid, tags=(tag,))

    def refresh_tree_row(self, idx: int):
        if self.df_workouts is None:
            return
        row = self.df_workouts.iloc[idx]
        iid = str(idx)

        week = row.get("Week", "")
        date = self.format_date_for_display(row.get("Date", ""))
        session = row.get("Session", "")
        sport = row.get("Sport", "")
        desc = row.get("Description", "")
        sched = row.get("ScheduledDate", "")

        if pd.isna(sched) or str(sched).strip().lower() in ("nan", "<na>"):
            sched = ""

        vals = (week, date, session, sport, desc, sched)

        if iid in self.tree_workouts.get_children():
            self.tree_workouts.item(iid, values=vals)
            self.apply_workout_row_style(idx)


    # ---------- Parameters editing ----------

    def edit_selected_parameter(self):
        if self.df_parameters is None or self.df_parameters.empty:
            messagebox.showwarning("Attenzione", "Nessun parametro caricato.")
            return
        sel = self.tree_params.selection()
        if not sel:
            messagebox.showwarning("Attenzione", "Seleziona un parametro.")
            return
        idx = int(sel[0])
        row = self.df_parameters.iloc[idx]

        win = tk.Toplevel(self.root)
        win.title("Modifica parametro")

        ttk.Label(win, text="Key:").grid(row=0, column=0, sticky="w")
        e_key = ttk.Entry(win)
        e_key.grid(row=0, column=1, sticky="we")
        e_key.insert(0, str(row.get("Key", "")))

        ttk.Label(win, text="Metric:").grid(row=1, column=0, sticky="w")
        e_metric = ttk.Entry(win)
        e_metric.grid(row=1, column=1, sticky="we")
        e_metric.insert(0, str(row.get("Metric", "")))

        ttk.Label(win, text="Expression:").grid(row=2, column=0, sticky="w")
        e_expr = ttk.Entry(win)
        e_expr.grid(row=2, column=1, sticky="we")
        e_expr.insert(0, str(row.get("Expression", "")))

        ttk.Label(win, text="Notes:").grid(row=3, column=0, sticky="w")
        e_notes = ttk.Entry(win)
        e_notes.grid(row=3, column=1, sticky="we")
        e_notes.insert(0, str(row.get("Notes", "")))

        win.columnconfigure(1, weight=1)

        def save_param():
            self.df_parameters.at[idx, "Key"] = e_key.get().strip()
            self.df_parameters.at[idx, "Metric"] = e_metric.get().strip()
            self.df_parameters.at[idx, "Expression"] = e_expr.get().strip()
            self.df_parameters.at[idx, "Notes"] = e_notes.get().strip()
            self.populate_params_tree()
            self.autosave_to_loaded_excel()
            win.destroy()

        ttk.Button(win, text="Salva", command=save_param).grid(row=4, column=0, columnspan=2, pady=5)

    # ---------- Garmin login ----------

    def login_garmin_with_credentials(self):
        email = self.entry_email.get().strip()
        password = self.entry_password.get().strip()
        if not email or not password:
            messagebox.showerror("Errore", "Inserisci email e password Garmin.")
            return
        try:
            self.garmin.login_with_credentials(email, password)
            self.lbl_garmin_status.config(text="● Connesso", foreground="green")
            messagebox.showinfo("OK", "Login a Garmin Connect effettuato (sessione salvata).")
        except Exception as e:
            self.lbl_garmin_status.config(text="● Non connesso", foreground="red")
            messagebox.showerror("Errore", f"Errore nel login a Garmin:\n{e}")

    def login_garmin_with_session(self):
        try:
            self.garmin.login_with_saved_session()
            self.lbl_garmin_status.config(text="● Connesso", foreground="green")
            messagebox.showinfo("OK", "Login a Garmin Connect effettuato tramite sessione salvata.")
        except Exception as e:
            self.lbl_garmin_status.config(text="● Non connesso", foreground="red")
            messagebox.showerror(
                "Errore",
                "Errore nel login con sessione salvata:\n"
                f"{e}\n\nSe necessario effettua un nuovo login con credenziali.",
            )

    # ---------- Garmin: upload / schedule / unschedule ----------

    def _get_selected_indices(self) -> list[int]:
        if self.df_workouts is None:
            messagebox.showerror("Errore", "Nessun Excel caricato.")
            return []
        sel = self.tree_workouts.selection()
        if not sel:
            messagebox.showerror("Errore", "Seleziona almeno un allenamento nella tabella.")
            return []
        return [int(iid) for iid in sel]

    def _run_with_loading(self, operation_func, title="Operazione in corso..."):
        """
        Esegue una funzione mostrando un dialog di loading.
        La funzione viene eseguita in un thread separato per non bloccare la GUI.
        
        Args:
            operation_func: Funzione da eseguire (deve ritornare (success, message))
            title: Titolo del dialog di loading
        """
        loading = LoadingDialog(self.root, title)
        result = {"success": False, "message": ""}
        
        def worker():
            try:
                success, message = operation_func(loading)
                result["success"] = success
                result["message"] = message
            except Exception as e:
                result["success"] = False
                result["message"] = str(e)
            finally:
                # Chiudi il loading nella GUI thread
                self.root.after(0, loading.close)
        
        # Avvia il worker thread
        thread = threading.Thread(target=worker, daemon=True)
        thread.start()
        
        # Aspetta che il thread finisca
        while thread.is_alive():
            self.root.update()
            thread.join(timeout=0.1)
        
        # Mostra il risultato
        if result["success"]:
            messagebox.showinfo("OK", result["message"])
        else:
            messagebox.showerror("Errore", result["message"])

    def upload_selected_workouts(self):
        if self.garmin.client is None:
            messagebox.showerror("Errore", "Non sei connesso a Garmin.")
            return
        idxs = self._get_selected_indices()
        if not idxs:
            return

        def operation(loading):
            created_ids = []
            total = len(idxs)
            
            for i, idx in enumerate(idxs, 1):
                loading.update_message(f"Caricamento workout {i}/{total}...")
                
                row = self.df_workouts.iloc[idx]
                steps_text = str(row.get("Steps", ""))
                workout_data = build_garmin_workout_from_excel_row(
                    row,
                    steps_text,
                    self.df_parameters,
                    use_prefix=self.use_prefix_var.get(),
                )

                resp = self.garmin.save_workout(workout_data)
                workout_id = resp.get("workoutId") or resp.get("workout_id")
                if workout_id:
                    self.df_workouts.at[idx, "WorkoutId"] = str(workout_id)
                    created_ids.append(workout_id)
                    self.root.after(0, lambda idx=idx: self.refresh_tree_row(idx))

            self.root.after(0, self.autosave_to_loaded_excel)
            return True, f"Creati {len(created_ids)} workout su Garmin."
        
        self._run_with_loading(operation, "Caricamento workout su Garmin")

    def normalize_workout_id(self, value) -> str:
        """Converte il WorkoutId in stringa pulita per le API Garmin."""
        import math

        if value is None:
            return ""

        if isinstance(value, float) and math.isnan(value):
            return ""

        s = str(value).strip()
        if s.lower() in ("nan", "<na>", ""):
            return ""
        if s.endswith(".0"):
            s = s[:-2]
        return s

    def upload_and_schedule_selected_workouts(self):
        """
        Crea (se serve) e pianifica su Garmin i workout selezionati,
        usando SEMPRE la data in colonna Date.
        Salva sia il WorkoutId che il WorkoutScheduleId restituito da Garmin.
        """
        if self.garmin.client is None:
            messagebox.showerror("Errore", "Non sei connesso a Garmin.")
            return

        idxs = self._get_selected_indices()
        if not idxs:
            return

        def operation(loading):
            total = 0
            num_workouts = len(idxs)

            for i, idx in enumerate(idxs, 1):
                loading.update_message(f"Elaborazione workout {i}/{num_workouts}...")
                
                row = self.df_workouts.iloc[idx]
                steps_text = str(row.get("Steps", ""))

                # ---- DATA: sempre quella della colonna Date ----
                date_val = row.get("Date", "")
                if isinstance(date_val, str):
                    date_str = date_val.strip()
                else:
                    try:
                        date_str = pd.to_datetime(date_val).date().isoformat()
                    except Exception:
                        raise ValueError(f"Data non valida per workout indice {idx}.")

                # Validazione formato YYYY-MM-DD
                try:
                    datetime.strptime(date_str, "%Y-%m-%d")
                except ValueError:
                    raise ValueError(f"Formato data non valido per workout indice {idx}: {date_str}")

                # ---- WORKOUT ID: se manca, crea il workout su Garmin ----
                workout_id = self.normalize_workout_id(row.get("WorkoutId", ""))

                if not workout_id:
                    loading.update_message(f"Creazione workout {i}/{num_workouts}...")
                    
                    workout_data = build_garmin_workout_from_excel_row(
                        row,
                        steps_text,
                        self.df_parameters,
                        use_prefix=self.use_prefix_var.get(),
                    )

                    resp = self.garmin.save_workout(workout_data)
                    workout_id = str(
                        (resp.get("workoutId") if isinstance(resp, dict) else "")
                        or (resp.get("workout_id") if isinstance(resp, dict) else "")
                    ).strip()

                    if not workout_id:
                        print(f"âš ï¸ Workout creato (riga {idx}) ma nessun 'workoutId' nella risposta.")
                        continue

                    self.df_workouts.at[idx, "WorkoutId"] = workout_id
                    self.root.after(0, lambda idx=idx: self.refresh_tree_row(idx))

                # ---- PIANIFICAZIONE: usa schedule_workout e salva WorkoutScheduleId ----
                loading.update_message(f"Pianificazione workout {i}/{num_workouts}...")
                
                resp_sched = self.garmin.schedule_workout(workout_id, date_str)

                schedule_id = ""
                if isinstance(resp_sched, dict):
                    schedule_id = str(
                        resp_sched.get("workoutScheduleId")
                        or resp_sched.get("workout_schedule_id")
                        or resp_sched.get("eventId")
                        or resp_sched.get("id")
                        or ""
                    ).strip()

                if schedule_id:
                    self.df_workouts.at[idx, "WorkoutScheduleId"] = schedule_id
                else:
                    print("âš ï¸ Nessun workoutScheduleId nella risposta:", resp_sched)

                self.df_workouts.at[idx, "ScheduledDate"] = date_str
                self.root.after(0, lambda idx=idx: self.refresh_tree_row(idx))
                total += 1

            self.root.after(0, self.autosave_to_loaded_excel)
            return True, f"Pianificati {total} workout."
        
        self._run_with_loading(operation, "Caricamento e pianificazione workout")


    def unschedule_selected_workouts(self):
        """Rimuove la pianificazione su Garmin usando il WorkoutScheduleId salvato."""
        if self.garmin.client is None:
            messagebox.showerror("Errore", "Non sei connesso a Garmin.")
            return

        idxs = self._get_selected_indices()
        if not idxs:
            return

        def operation(loading):
            removed = 0
            num_workouts = len(idxs)

            for i, idx in enumerate(idxs, 1):
                loading.update_message(f"Rimozione pianificazione {i}/{num_workouts}...")
                
                row = self.df_workouts.iloc[idx]

                # id della pianificazione, NON il workoutId
                schedule_id = self.normalize_workout_id(row.get("WorkoutScheduleId", ""))
                sched_date = row.get("ScheduledDate", "")
                sched_date = "" if pd.isna(sched_date) else str(sched_date).strip()
                if sched_date.lower() in ("nan", "<na>"):
                    sched_date = ""

                # Se manca uno dei due, non posso fare nulla
                if not schedule_id or not sched_date:
                    continue

                try:
                    self.garmin.unschedule_workout(schedule_id, sched_date)
                    # pulisco sia la data che l'id di pianificazione
                    self.df_workouts.at[idx, "ScheduledDate"] = ""
                    self.df_workouts.at[idx, "WorkoutScheduleId"] = ""
                    self.root.after(0, lambda idx=idx: self.refresh_tree_row(idx))
                    removed += 1
                except Exception as e:
                    msg = str(e)
                    if "403" in msg:
                        error_msg = (
                            f"Garmin ha risposto 403 Forbidden nel tentativo di rimuovere "
                            f"la pianificazione (scheduleId {schedule_id}, {sched_date}).\n\n"
                            f"Questo significa che l'API usata non Ã¨ autorizzata "
                            f"a cancellare la programmazione. Per questo allenamento dovrai "
                            f"rimuovere la pianificazione manualmente da Garmin Connect.\n\n"
                            f"Dettagli tecnici:\n{msg}"
                        )
                        raise RuntimeError(error_msg)
                    else:
                        error_msg = f"Errore nel rimuovere la pianificazione (scheduleId {schedule_id}, {sched_date}):\n{msg}"
                        raise RuntimeError(error_msg)

            if removed:
                self.root.after(0, self.autosave_to_loaded_excel)
                return True, f"Rimossi {removed} workout dalla programmazione."
            else:
                msg = ("Nessun workout Ã¨ stato rimosso.\n"
                       "Verifica che le righe selezionate abbiano sia ScheduledDate sia WorkoutScheduleId compilati.")
                return True, msg
        
        self._run_with_loading(operation, "Rimozione pianificazione workout")



    def delete_workouts_from_garmin(self):
        """
        Per ogni riga selezionata:
        - se ha una pianificazione (WorkoutScheduleId + ScheduledDate), prova a rimuoverla;
        - cancella il workout dalla libreria Garmin;
        - pulisce WorkoutId, WorkoutScheduleId e ScheduledDate nell'Excel.
        """
        if self.garmin.client is None:
            messagebox.showerror("Errore", "Non sei connesso a Garmin.")
            return

        idxs = self._get_selected_indices()
        if not idxs:
            return

        # Conferma utente (operazione distruttiva)
        confirm_msg = (
            "Vuoi cancellare DEFINITIVAMENTE i workout selezionati dalla libreria Garmin?\n"
            "Se sono pianificati, verrÃ  prima rimossa la pianificazione."
        )
        if messagebox.askyesno("Conferma", confirm_msg) is False:
            return

        def operation(loading):
            deleted = 0
            num_workouts = len(idxs)

            for i, idx in enumerate(idxs, 1):
                loading.update_message(f"Cancellazione workout {i}/{num_workouts}...")
                
                row = self.df_workouts.iloc[idx]

                workout_id = self.normalize_workout_id(row.get("WorkoutId", ""))
                if not workout_id:
                    # niente da cancellare su Garmin
                    continue

                schedule_id = self.normalize_workout_id(row.get("WorkoutScheduleId", ""))
                sched_date = row.get("ScheduledDate", "")
                sched_date = "" if pd.isna(sched_date) else str(sched_date).strip()
                if sched_date.lower() in ("nan", "<na>"):
                    sched_date = ""

                # 1) se c'Ã¨ una pianificazione, prova a toglierla
                if schedule_id and sched_date:
                    try:
                        loading.update_message(f"Rimozione pianificazione workout {i}/{num_workouts}...")
                        self.garmin.unschedule_workout(schedule_id, sched_date)
                        print(
                            f"Dis-pianificato workoutScheduleId {schedule_id} ({sched_date}) "
                            f"prima della cancellazione definitiva."
                        )
                    except Exception as e:
                        # Non blocco la cancellazione del workout, ma avviso
                        print(
                            f"âš ï¸ Errore nel rimuovere la pianificazione (scheduleId {schedule_id}): {e}"
                        )

                # 2) cancella il workout dalla libreria
                loading.update_message(f"Cancellazione workout {i}/{num_workouts}...")
                self.garmin.delete_workout(workout_id)

                # 3) pulisco i campi collegati alla parte Garmin
                self.df_workouts.at[idx, "WorkoutId"] = ""
                if "WorkoutScheduleId" in self.df_workouts.columns:
                    self.df_workouts.at[idx, "WorkoutScheduleId"] = ""
                if "ScheduledDate" in self.df_workouts.columns:
                    self.df_workouts.at[idx, "ScheduledDate"] = ""
                self.root.after(0, lambda idx=idx: self.refresh_tree_row(idx))
                deleted += 1

            if deleted:
                self.root.after(0, self.autosave_to_loaded_excel)
                return True, f"Cancellati definitivamente {deleted} workout dalla libreria Garmin."
            else:
                return True, "Nessun workout Ã¨ stato cancellato (nessun WorkoutId valido nelle righe selezionate)."
        
        self._run_with_loading(operation, "Cancellazione workout da Garmin")



def main():
    root = tk.Tk()
    app = TrainingPlannerGUI(root)
    root.mainloop()


if __name__ == "__main__":
    main()